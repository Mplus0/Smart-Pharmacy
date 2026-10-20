"""Dependency-isolated contract tests, NOT ROS/hardware integration validation."""
import hashlib
import importlib
import json
import math
import sqlite3
import sys
import tempfile
import threading
import types
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch


class Message:
    def __init__(self, data=None, **kwargs):
        self.data = data
        self.__dict__.update(kwargs)


class Publisher:
    def __init__(self):
        self.messages = []

    def publish(self, msg):
        self.messages.append(msg.data if hasattr(msg, "data") else msg)


def module(name, **attributes):
    result = types.ModuleType(name)
    result.__dict__.update(attributes)
    return result


class OfflineContracts(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.ros = MagicMock()
        cls.ros.is_shutdown.return_value = False
        cls.ros.get_param.side_effect = lambda name, default: default
        cls.ros.ROSInterruptException = RuntimeError
        stubs = {"rospy": cls.ros, "cv_bridge": module("cv_bridge"),
                 "cv2": module("cv2"), "numpy": module("numpy"),
                 "actionlib": module("actionlib"),
                 "pyzbar": module("pyzbar"), "pyzbar.pyzbar": module("pyzbar.pyzbar", decode=lambda frame: [])}
        for package, submodule, names in [
            ("std_msgs", "msg", ["Int32MultiArray", "Int32", "String"]),
            ("sensor_msgs", "msg", ["Image", "CompressedImage"]),
            ("geometry_msgs", "msg", ["Pose", "Point", "Quaternion", "Twist"]),
            ("nav_msgs", "msg", ["Odometry"]),
            ("move_base_msgs", "msg", ["MoveBaseGoal", "MoveBaseAction"]),
            ("std_srvs", "srv", ["Empty", "Trigger", "TriggerResponse"]),
        ]:
            stubs[package] = module(package)
            stubs[package + "." + submodule] = module(package + "." + submodule,
                                                       **{name: Message for name in names})
        stubs["actionlib_msgs"] = module("actionlib_msgs")
        stubs["actionlib_msgs.msg"] = module("actionlib_msgs.msg", GoalStatus=types.SimpleNamespace(SUCCEEDED=3))
        stubs["tf"] = module("tf")
        stubs["tf.transformations"] = module("tf.transformations",
            quaternion_from_euler=lambda x, y, z, **kw: [0, 0, math.sin(z / 2), math.cos(z / 2)])
        cls.saved = set(sys.modules)
        cls.context = patch.dict(sys.modules, stubs)
        cls.context.start()
        cls.runtime = importlib.import_module("pharmacy_iscc.perception.runtime")
        cls.nav = importlib.import_module("pharmacy_iscc.navigation.state_machine")
        cls.dual = importlib.import_module("pharmacy_iscc.communication.dual_car")
        cls.referee = importlib.import_module("pharmacy_iscc.communication.referee")
        cls.board2 = importlib.import_module("pharmacy_iscc.communication.board2")
        cls.records = importlib.import_module("pharmacy_iscc.recording.records")

    @classmethod
    def tearDownClass(cls):
        cls.context.stop()

    def setUp(self):
        self.ros.reset_mock()
        self.ros.is_shutdown.return_value = False
        self.ros.get_param.side_effect = lambda name, default: default
        self.runtime.nav_state_cb(Message(9))
        self.runtime.nav_state_cb(Message(10))
        self.runtime.pub_flag = Publisher()
        self.runtime.pub_board2 = Publisher()
        self.runtime.pub_board1_all_text = Publisher()

    def test_stability_batch_layout_and_one_publish_per_visit(self):
        r = self.runtime
        text = ["D", "", "ABC", "BC"]
        with patch.object(r, "play_board1_error_audio") as audio, patch.object(r, "record_result"):
            for _ in range(2):
                self.assertIsNone(r.stable_publish_board1(text))
            self.assertEqual(r.pub_flag.messages, [])
            result = r.stable_publish_board1(text)
            self.assertEqual(result, [1, 1, 1, 3, 2, 1, 1, 0, 1, 2, 3, 1])
            r.stable_publish_board1(text)
            self.assertEqual(len(r.pub_flag.messages), 1)
            audio.assert_called_once()
            shared = json.loads(r.pub_board1_all_text.messages[0])
            self.assertEqual(shared["all_text"], text)
            self.assertEqual(shared["selected_msg"], result[:6])
            self.assertEqual(len(r.pub_board1_all_text.messages), 3)

    def test_invalid_observation_resets_stability(self):
        r = self.runtime
        for text in [["ABC", "", "", ""]] * 2 + [["", "", "", ""]] + [["ABC", "", "", ""]] * 2:
            r.stable_publish_board1(text)
        self.assertEqual(r.pub_flag.messages, [])
        r.stable_publish_board1(["ABC", "", "", ""])
        self.assertEqual(len(r.pub_flag.messages), 1)

    def test_board2_get_protocol_and_cancellation(self):
        for wait in [0, 5, 6, 7, 8, 9, 10, True, 4, 11, "5", None]:
            conn = MagicMock()
            conn.recv.return_value = (json.dumps({"ok": True, "wait_time": wait}) + "\n").encode()
            with patch.object(self.board2.socket, "create_connection", return_value=conn):
                result = self.board2.query_board2_tcp(lambda: False)
            valid = type(wait) is int and (wait == 0 or 5 <= wait <= 10)
            self.assertEqual(result, ("free" if wait == 0 else "busy_%d" % wait) if valid else None)
            conn.sendall.assert_called_once_with(b"GET\n")
            conn.close.assert_called_once()
        conn = MagicMock()
        with patch.object(self.board2.socket, "create_connection", return_value=conn):
            self.assertIsNone(self.board2.query_board2_tcp(lambda: True))
        conn.recv.assert_not_called()
        conn.close.assert_called_once()

    def test_board2_default_free_and_visit_guard(self):
        r = self.runtime
        r.nav_state_cb(Message(13))
        with patch.object(r, "query_board2_tcp", return_value=None):
            r.publish_board2_tcp_result()
            r.publish_board2_tcp_result()
        self.assertEqual(r.pub_board2.messages, [[0, 0]])
        r.nav_state_cb(Message(12))
        r.nav_state_cb(Message(13))
        def changed_visit(cancelled):
            r.nav_state_cb(Message(12))
            r.nav_state_cb(Message(13))
            return "busy_5"
        with patch.object(r, "query_board2_tcp", side_effect=changed_visit):
            r.publish_board2_tcp_result()
        self.assertEqual(r.pub_board2.messages, [[0, 0]])

    def test_single_car_batch_and_dual_car_first_task(self):
        for dual in [False, True]:
            self.ros.get_param.side_effect = lambda name, default: dual if name == "~enable_dual_car" else default
            app = self.nav.MoveBaseSquare.__new__(self.nav.MoveBaseSquare)
            app.init_runtime_state()
            app.count = 10
            app.recognition_records = MagicMock()
            app.pub_referee_cv2 = Publisher()
            data = [1, 1, 1, 3, 2, 1, 1, 0, 1, 2, 3, 1]
            app.detect_board1_result(Message(data))
            self.assertEqual(app.pending_board1_tasks, [] if dual else [data[6:]])
            self.assertEqual(app.cv2_result, "ABC-3")
            app.recognition_records.record.assert_called_once()
            app.detect_board1_result(Message(data))
            app.recognition_records.record.assert_called_once()

    def test_dual_car_done_dedup_and_peer_result_selection(self):
        self.ros.get_param.side_effect = lambda name, default: True if name == "~enable_dual_car" else default
        app = self.nav.MoveBaseSquare.__new__(self.nav.MoveBaseSquare)
        app.init_runtime_state()
        app.pub_nav_state = Publisher()
        app.pub_round_done = Publisher()
        app.release_next_car_if_needed()
        app.release_next_car_if_needed()
        self.assertEqual(app.pub_round_done.messages, [[1, 1]])
        self.assertFalse(app.have_turn)
        app.count = 8
        app.peer_done_callback(Message([2, 1]))
        self.assertTrue(app.have_turn)
        self.assertEqual(app.count, 9)
        app.peer_done_callback(Message([2, 1]))
        self.assertEqual(app.pub_nav_state.messages, [9])
        app.use_peer_board1_result = True
        app.peer_board1_data = {"car_id": 2, "seq": 1, "selected_index": 2,
                               "all_text": ["A", "", "ABC", "BC"]}
        app.apply_board1_selected_msg = MagicMock()
        self.assertTrue(app.try_use_peer_board1_result_before_board1_navigation())
        app.apply_board1_selected_msg.assert_called_once_with([1, 0, 1, 2, 3, 0])
        self.assertFalse(app.try_use_peer_board1_result_before_board1_navigation())

    def test_dual_tcp_wire_dedup_peer_ip_and_token(self):
        app = self.dual.DualCarTcpLink.__new__(self.dual.DualCarTcpLink)
        app.last_peer_seq = 0
        app.publish_peer_done = MagicMock()
        line = json.dumps({"type": "round_done", "car_id": 2, "seq": 1})
        app.handle_line(line, ("wrong", 1))
        app.publish_peer_done.assert_not_called()
        app.handle_line(line, (self.dual.PEER_IP, 1))
        app.handle_line(line, (self.dual.PEER_IP, 1))
        app.publish_peer_done.assert_called_once_with(2, 1)
        sock = MagicMock()
        with patch.object(self.dual.socket, "socket", return_value=sock), patch.object(self.dual, "SHARED_TOKEN", "secret"):
            app.send_done_once(1, 2)
        payload = sock.sendall.call_args.args[0]
        self.assertTrue(payload.endswith(b"\n"))
        self.assertEqual(json.loads(payload), {"type": "round_done", "car_id": 1, "seq": 2, "token": "secret"})

    def test_referee_field_semantics(self):
        app = self.referee.RefereeClient.__new__(self.referee.RefereeClient)
        app.payload_lock = threading.Lock()
        app.payload = app.default_payload()
        self.assertEqual(app.payload, {"id": "1", "speed": 0.0, "odom": [0.0, 0.0],
                                       "task": "R", "CV1": "None", "CV2": "None"})
        app.task_cb(Message("invalid"))
        self.assertEqual(app.payload["task"], "R")
        app.task_cb(Message("A"))
        app.cv1_cb(Message("WAIT-5"))
        app.cv2_cb(Message("ABC-3"))
        self.assertEqual((app.payload["task"], app.payload["CV1"], app.payload["CV2"]),
                         ("A", "WAIT-5", "ABC-3"))

    def test_records_digest_export_once_and_late_result_ignored(self):
        from openpyxl import load_workbook
        with tempfile.TemporaryDirectory() as folder, patch.object(self.records, "LOG_DIR", folder):
            task = self.records.RecognitionRecords(1)
            value = {"状态": "忙碌", "等待秒数": 5}
            task.record("识别板二", "识别结果", value)
            path = task.export()
            self.assertEqual(task.export(), path)
            task.record("识别板二", "识别结果", {"等待秒数": 10})
            with sqlite3.connect(self.records.database_path(1)) as connection:
                rows = connection.execute("SELECT content, md5 FROM records").fetchall()
                self.assertEqual(connection.execute("SELECT count(*) FROM active_run").fetchone()[0], 0)
            self.assertEqual(len(rows), 1)
            self.assertEqual(rows[0][0], json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":")))
            self.assertEqual(rows[0][1], hashlib.md5(rows[0][0].encode()).hexdigest())
            workbook = load_workbook(path)
            self.assertEqual(workbook.active.max_row, 2)
            self.assertEqual(workbook.active["D2"].value, rows[0][0])
            self.assertEqual(workbook.active["E2"].value, rows[0][1])
            workbook.close()
            self.records.RecognitionRecords(1)
            with sqlite3.connect(self.records.database_path(1)) as connection:
                self.assertEqual(connection.execute("SELECT count(*) FROM records").fetchone()[0], 0)
